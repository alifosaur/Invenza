from typing import List, Optional

import io
from fastapi import UploadFile, File, Form, BackgroundTasks
from fastapi.responses import StreamingResponse
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from app.models.location import Location
from app.models.warehouse import Warehouse
from pydantic import BaseModel
from fastapi import APIRouter, Depends, HTTPException, status, Query
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.db.session import get_db
from app.models.product import Product, Category
from app.models.stock import Stock
from app.schemas.product import ProductCreate, ProductUpdate, ProductOut, CategoryCreate, CategoryOut
from app.api.deps import get_current_user, require_manager
from app.models.user import User

router = APIRouter(tags=["products"])


# ── Categories ─────────────────────────────────────────────────────────────

@router.get("/categories", response_model=List[CategoryOut])
async def list_categories(db: AsyncSession = Depends(get_db), _: User = Depends(get_current_user)):
    result = await db.execute(select(Category))
    return result.scalars().all()


@router.post("/categories", response_model=CategoryOut, status_code=status.HTTP_201_CREATED)
async def create_category(
    payload: CategoryCreate,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_manager),
):
    cat = Category(name=payload.name)
    db.add(cat)
    await db.flush()
    await db.refresh(cat)
    return cat


# ── Products ───────────────────────────────────────────────────────────────

@router.get("/products", response_model=List[ProductOut])
async def list_products(
    search: Optional[str] = Query(None),
    category_id: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
):
    query = select(Product).options(selectinload(Product.category), selectinload(Product.stock_entries))
    if search:
        query = query.where(
            Product.name.ilike(f"%{search}%") | Product.sku.ilike(f"%{search}%")
        )
    if category_id:
        query = query.where(Product.category_id == category_id)
    result = await db.execute(query)
    return result.scalars().all()


@router.post("/products", response_model=ProductOut, status_code=status.HTTP_201_CREATED)
async def create_product(
    payload: ProductCreate,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_manager),
):
    # Check SKU uniqueness
    existing = await db.execute(select(Product).where(Product.sku == payload.sku))
    if existing.scalar_one_or_none():
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="SKU already exists")

    product = Product(
        name=payload.name,
        sku=payload.sku,
        category_id=payload.category_id,
        uom=payload.uom,
        reorder_threshold=payload.reorder_threshold,
        image_data=payload.image_data,
    )
    db.add(product)
    await db.flush()

    # Seed initial stock if provided
    if payload.initial_stock and payload.initial_location_id:
        stock = Stock(
            product_id=product.id,
            location_id=payload.initial_location_id,
            on_hand=payload.initial_stock,
            per_unit_cost=payload.per_unit_cost or 0,
        )
        db.add(stock)

    await db.flush()
    await db.refresh(product)
    return product


@router.get("/products/{product_id}", response_model=ProductOut)
async def get_product(
    product_id: str,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
):
    result = await db.execute(
        select(Product).options(selectinload(Product.category), selectinload(Product.stock_entries)).where(Product.id == product_id)
    )
    product = result.scalar_one_or_none()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    return product


@router.put("/products/{product_id}", response_model=ProductOut)
async def update_product(
    product_id: str,
    payload: ProductUpdate,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_manager),
):
    result = await db.execute(select(Product).where(Product.id == product_id))
    product = result.scalar_one_or_none()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    for field, value in payload.model_dump(exclude_none=True).items():
        setattr(product, field, value)
    await db.flush()
    await db.refresh(product)
    return product


@router.delete("/products/{product_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_product(
    product_id: str,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_manager),
):
    result = await db.execute(select(Product).where(Product.id == product_id))
    product = result.scalar_one_or_none()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    await db.delete(product)


# ── Excel Import / Export ──────────────────────────────────────────────────

@router.get("/products/template")
async def get_excel_template(_: User = Depends(get_current_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"

    headers = ["SKU", "Product Name", "Category", "Unit of Measure", "Quantity", "Unit Price", "Reorder Threshold", "Image URL"]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    # Add Instructions sheet
    ws_inst = wb.create_sheet("Instructions")
    ws_inst.append(["Column", "Required", "Description"])
    ws_inst.append(["SKU", "Yes", "Unique identifier for the product"])
    ws_inst.append(["Product Name", "Yes", "Name of the product"])
    ws_inst.append(["Category", "No", "Product category (will be created if missing)"])
    ws_inst.append(["Unit of Measure", "No", "e.g., pcs, kg, ml"])
    ws_inst.append(["Quantity", "No", "Initial stock quantity (numeric)"])
    ws_inst.append(["Unit Price", "No", "Price per unit (numeric)"])
    ws_inst.append(["Reorder Threshold", "No", "Minimum stock level for alerts (numeric)"])
    ws_inst.append(["Image URL", "No", "Public URL to product image"])

    # Make bold
    for cell in ws_inst[1]:
        cell.font = Font(bold=True)
        
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="invenza_inventory_template.xlsx"'}
    )


@router.get("/products/export/excel")
async def export_products_excel(
    search: Optional[str] = Query(None),
    category_id: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
):
    query = select(Product).options(selectinload(Product.category), selectinload(Product.stock_entries))
    if search:
        query = query.where(
            Product.name.ilike(f"%{search}%") | Product.sku.ilike(f"%{search}%")
        )
    if category_id:
        query = query.where(Product.category_id == category_id)
        
    result = await db.execute(query)
    products = result.scalars().all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    
    headers = ["SKU", "Product Name", "Category", "Quantity", "Unit Price", "Total Value", "Reorder Threshold", "Image URL", "Created At"]
    ws.append(headers)
    
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        
    total_qty = 0
    total_val = 0
    
    for p in products:
        qty = sum(s.on_hand for s in p.stock_entries)
        price = p.stock_entries[0].per_unit_cost if p.stock_entries else 0
        val = qty * price
        
        total_qty += qty
        total_val += val
        
        ws.append([
            p.sku,
            p.name,
            p.category.name if p.category else "",
            qty,
            price,
            val,
            p.reorder_threshold,
            p.image_data if (p.image_data and str(p.image_data).startswith('http')) else "",
            p.created_at.strftime("%Y-%m-%d %H:%M")
        ])
        
    ws_summary = wb.create_sheet("Summary")
    ws_summary.append(["Metric", "Value"])
    ws_summary.append(["Total Products", len(products)])
    ws_summary.append(["Total Quantity", total_qty])
    ws_summary.append(["Total Inventory Value", total_val])
    ws_summary.append(["Export Date", datetime.now().strftime("%Y-%m-%d %H:%M")])
    for cell in ws_summary[1]:
        cell.font = Font(bold=True)
        
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    
    date_str = datetime.now().strftime("%Y-%m-%d")
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="warehouse_inventory_{date_str}.xlsx"'}
    )


class ImportResponse(BaseModel):
    valid_count: int
    warning_count: int
    error_count: int
    rows: list
    
@router.post("/products/import/excel")
async def import_products_excel(
    file: UploadFile = File(...),
    confirm: bool = Form(False),
    update_existing: bool = Form(True),
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_manager),
):
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="Invalid file type. Please upload an Excel file.")
        
    content = await file.read()
    try:
        wb = load_workbook(filename=io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not parse Excel file.")
        
    ws = wb.active
    
    rows_data = []
    headers = []
    
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(h).strip().lower() if h else "" for h in row]
            continue
            
        row_dict = dict(zip(headers, row))
        rows_data.append(row_dict)
        
    # Validation phase
    results = []
    valid_count = 0
    warning_count = 0
    error_count = 0
    
    # Pre-fetch existing SKUs and categories for quick validation
    existing_products_result = await db.execute(select(Product))
    existing_products = {p.sku: p for p in existing_products_result.scalars().all()}
    
    existing_categories_result = await db.execute(select(Category))
    existing_categories = {c.name.lower(): c for c in existing_categories_result.scalars().all()}
    
    # Pre-fetch a default location for initial stock if provided
    location_result = await db.execute(select(Location).limit(1))
    default_location = location_result.scalar_one_or_none()
    
    new_products_to_create = []
    products_to_update = []
    new_stocks_to_create = []
    categories_to_create = set()
    
    seen_skus = set()
    
    for r_idx, row in enumerate(rows_data, start=2):
        status = "Valid"
        issues = []
        is_error = False
        
        sku = row.get("sku")
        name = row.get("product name")
        cat_name = row.get("category")
        uom = row.get("unit of measure")
        qty = row.get("quantity")
        price = row.get("unit price")
        reorder = row.get("reorder threshold")
        image_url = row.get("image url")
        
        if not sku:
            issues.append("Missing SKU")
            is_error = True
        elif sku in seen_skus:
            issues.append("Duplicate SKU in file")
            is_error = True
        else:
            seen_skus.add(sku)
            
        if not name:
            issues.append("Missing Product Name")
            is_error = True
            
        try:
            if qty is not None:
                qty = float(qty)
                if qty < 0:
                    issues.append("Invalid quantity (negative)")
                    is_error = True
        except ValueError:
            issues.append("Quantity must be numeric")
            is_error = True
            
        try:
            if price is not None:
                price = float(price)
                if price < 0:
                    issues.append("Invalid price (negative)")
                    is_error = True
        except ValueError:
            issues.append("Unit Price must be numeric")
            is_error = True
            
        if not is_error:
            if sku in existing_products:
                if not update_existing:
                    issues.append("SKU exists (Skipping)")
                    status = "Warning"
                    warning_count += 1
                else:
                    issues.append("SKU exists (Will Update)")
                    status = "Warning"
                    warning_count += 1
                    products_to_update.append({"product": existing_products[sku], "row": row})
            else:
                valid_count += 1
                new_products_to_create.append(row)
                
            if cat_name and str(cat_name).lower() not in existing_categories:
                categories_to_create.add(str(cat_name))
        else:
            error_count += 1
            status = "Error"
            
        results.append({
            "row_num": r_idx,
            "sku": sku or "",
            "name": name or "",
            "status": status,
            "issues": issues
        })
        
    if not confirm:
        return {
            "valid_count": valid_count,
            "warning_count": warning_count,
            "error_count": error_count,
            "rows": results
        }
        
    if error_count > 0:
        raise HTTPException(status_code=400, detail="Cannot import file with critical errors. Please fix and re-upload.")
        
    # Transactional Import
    try:
        # Create missing categories
        for cname in categories_to_create:
            if cname.lower() not in existing_categories:
                new_cat = Category(name=cname)
                db.add(new_cat)
                await db.flush()
                existing_categories[cname.lower()] = new_cat
                
        # Update existing
        if update_existing:
            for item in products_to_update:
                p = item["product"]
                r = item["row"]
                if r.get("product name"): p.name = str(r.get("product name"))
                if r.get("unit of measure"): p.uom = str(r.get("unit of measure"))
                if r.get("reorder threshold") is not None: 
                    try: p.reorder_threshold = float(r.get("reorder threshold"))
                    except ValueError: pass
                if r.get("category"):
                    p.category_id = existing_categories[str(r.get("category")).lower()].id
                if r.get("image url"):
                    p.image_data = str(r.get("image url"))
                    
        # Create new
        for r in new_products_to_create:
            cat_id = None
            if r.get("category"):
                cat_id = existing_categories[str(r.get("category")).lower()].id
                
            p = Product(
                sku=str(r.get("sku")),
                name=str(r.get("product name")),
                uom=str(r.get("unit of measure") or "pcs"),
                category_id=cat_id,
            )
            if r.get("reorder threshold") is not None:
                try: p.reorder_threshold = float(r.get("reorder threshold"))
                except ValueError: pass
            if r.get("image url"):
                p.image_data = str(r.get("image url"))
                
            db.add(p)
            await db.flush()
            
            qty = r.get("quantity")
            price = r.get("unit price")
            if qty and float(qty) > 0 and default_location:
                stock = Stock(
                    product_id=p.id,
                    location_id=default_location.id,
                    on_hand=float(qty),
                    per_unit_cost=float(price) if price else 0
                )
                db.add(stock)
                
        await db.commit()
        return {"detail": f"Successfully imported {len(new_products_to_create)} new products and updated {len(products_to_update)}."}
        
    except Exception as e:
        await db.rollback()
        raise HTTPException(status_code=500, detail=f"Database transaction failed: {str(e)}")
